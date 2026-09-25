"""Pure regressions for the isolated backend audit fixes."""

from __future__ import annotations

import datetime as dt
import decimal
import io
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import HTTPException, UploadFile

from app.api import imports as imports_api
from app.documents.engine.branding import BrandingSnapshot
from app.documents.engine.context import InventoryDocContext
from app.documents.modules.inventory import audit_export


class TrackingBytesIO(io.BytesIO):
    def __init__(self, initial_bytes: bytes) -> None:
        super().__init__(initial_bytes)
        self.read_sizes: list[int] = []

    def read(self, size: int = -1) -> bytes:
        self.read_sizes.append(size)
        return super().read(size)


def test_xlsx_upload_reads_at_most_limit_plus_one(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        imports_api, "get_settings", lambda: SimpleNamespace(MAX_UPLOAD_MB=1)
    )
    file = TrackingBytesIO(b"x" * (1024 * 1024 + 2))
    upload = UploadFile(filename="large.xlsx", file=file)

    with pytest.raises(HTTPException) as raised:
        imports_api._read_xlsx(upload)

    assert raised.value.status_code == 413
    assert file.read_sizes == [1024 * 1024 + 1]


def test_audit_export_contains_persisted_detail_sheets_and_safe_text() -> None:
    now = dt.datetime(2026, 9, 24, 12, 30, tzinfo=dt.UTC)
    branding = BrandingSnapshot(
        legal_name="Inventario Dedalo",
        tax_id="20600000001",
        tax_id_label="RUC",
        address=None,
        phone=None,
        email=None,
        website=None,
        brand_color_primary="#1F4E79",
        brand_color_secondary="#F5F7FA",
        brand_color_accent="#C00000",
        currency_style="PEN",
        date_format="yyyy-mm-dd",
        footer_text=None,
        footer_left=None,
        footer_right=None,
        logo_media_type=None,
        logo_sha256="b" * 64,
        settings_version=3,
    )
    ctx = InventoryDocContext(
        branding=branding,
        campaign={"code": "CAM-1", "name": "Conteo", "status": "APPROVED"},
        valuation={"summary": {}, "items": [], "warnings": []},
        document_number="AUD-000001",
        title="Exportacion de Auditoria",
        document_type="audit-export",
        generated_at=now,
        generated_by_label="Auditor",
        sessions=[
            {
                "session_number": 1,
                "session_type": "PRIMARY",
                "status": "SUBMITTED",
                "started_at": now,
                "submitted_at": now,
                "device_identifier": "scanner-1",
                "user_display_name": "Operador Uno",
            }
        ],
        counts=[
            {
                "session_number": 1,
                "product_reference": "SKU-1",
                "product_name": "Producto Uno",
                "quantity": decimal.Decimal("3"),
                "damaged_quantity": decimal.Decimal("1"),
                "updated_at": now,
            }
        ],
        events=[
            {
                "session_number": 1,
                "server_sequence": 1,
                "event_type": "QR_SCAN",
                "product_reference": "SKU-1",
                "scanned_code": "SKU-1",
                "quantity": decimal.Decimal("1"),
                "damage_delta_quantity": decimal.Decimal("0"),
                "occurred_at": now,
            }
        ],
        damages=[
            {
                "session_number": 1,
                "product_reference": "SKU-1",
                "scanned_code": None,
                "action": "DAMAGE_ADD",
                "quantity": decimal.Decimal("1"),
                "reason": "Golpe",
                "observation": "Caja abierta",
                "event_id": "event-1",
                "created_by": "Operador Uno",
                "created_at": now,
                "has_evidence": True,
            }
        ],
        extras=[
            {
                "session_number": 1,
                "product_reference": "SKU-EXTRA",
                "product_name": "Producto Extra",
                "quantity": decimal.Decimal("2"),
                "first_detected_at": now,
            }
        ],
        unknowns=[
            {
                "session_number": 1,
                "scanned_code": "=1+1",
                "quantity": decimal.Decimal("1"),
                "damaged_quantity": decimal.Decimal("0"),
                "resolved_product_reference": None,
                "resolved_by": None,
                "resolved_at": None,
            }
        ],
        recounts=[
            {
                "status": "COMPLETED",
                "requested_by": "Supervisor",
                "assigned_to": "Operador Dos",
                "source_session_number": 1,
                "resulting_session_number": 2,
                "reason": "Validacion",
                "started_at": now,
                "completed_at": now,
                "cancelled_at": None,
                "cancel_reason": None,
            }
        ],
        audit_events=[
            {
                "action": "COUNT_SUBMITTED",
                "actor": "Operador Uno",
                "entity_type": "inventory_count_session",
                "entity_id": "session-1",
                "occurred_at": now,
                "metadata": '{"session_number":1}',
            }
        ],
        reconciliation_rows=[
            {
                "product": {"internal_reference": "SKU-1", "name": "Producto Uno"},
                "expected_quantity": decimal.Decimal("4"),
                "approved_physical_quantity": decimal.Decimal("3"),
                "difference_quantity": decimal.Decimal("-1"),
                "damaged_quantity": decimal.Decimal("1"),
                "missing_quantity": decimal.Decimal("1"),
                "surplus_quantity": decimal.Decimal("0"),
                "status": "APPROVED",
                "reason": "Faltante",
                "observation": "Revision",
                "effective_unit_cost": decimal.Decimal("2"),
                "missing_cost_value": decimal.Decimal("2"),
                "damage_cost_value": decimal.Decimal("2"),
                "surplus_cost_value": decimal.Decimal("0"),
                "affected_sale_value": decimal.Decimal("5"),
            }
        ],
        warnings=[],
    )

    data = audit_export.render(
        ctx,
        extra_metadata={"source_sha256": "a" * 64, "template_version": "1"},
    )
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)

    assert workbook.sheetnames == list(audit_export.SHEET_NAMES)
    assert workbook["Sesiones"]["G2"].value == "Operador Uno"
    assert workbook["Sesiones"]["D2"].value == now.replace(tzinfo=None)
    assert workbook["Conteos"]["B2"].value == "SKU-1"
    assert isinstance(workbook["Conteos"]["D2"].value, (int, float))
    assert workbook["Eventos"]["C2"].value == "QR_SCAN"
    assert isinstance(workbook["Eventos"]["H2"].value, dt.datetime)
    assert workbook["Danos"]["F2"].value == "Golpe"
    assert workbook["Extras"]["B2"].value == "SKU-EXTRA"
    assert workbook["Unknowns"]["B2"].value == "'=1+1"
    assert workbook["Reconteos"]["A2"].value == "COMPLETED"
    assert workbook["Auditoria"]["A2"].value == "COUNT_SUBMITTED"
    assert workbook["Conciliacion"]["A2"].value == "SKU-1"
    summary = {
        row[0]: row[1]
        for row in workbook["Resumen"].iter_rows(min_row=2, values_only=True)
        if row[0]
    }
    assert summary["source_sha256"] == "a" * 64
    assert summary["document_number"] == "AUD-000001"
