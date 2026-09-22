"""Pruebas unitarias del pipeline de importacion (parsers, mappings, cost)."""

from __future__ import annotations

import datetime as dt
import decimal

import openpyxl
import pytest

from app.importers import mappings
from app.importers.base import MissingHeaderError, UnknownImportTypeError
from app.importers.contacts import parse_contacts
from app.importers.cost import resolve_effective_cost
from app.importers.excel_reader import load_workbook
from app.importers.products import parse_mixed_product_export, parse_product_master
from app.importers.validators import parse_datetime_aware, parse_decimal
from app.models.enums import CostSource, ImportBatchType
from tests.import_helpers import make_xlsx

CONTACT_HEADERS = [
    "Actividades",
    "Apellido materno",
    "Apellido paterno",
    "Ciudad",
    "Correo electrónico",
    "Nombre completo",
    "Nombres",
    "País",
    "Teléfono",
    "Vendedor",
    "Etiquetas",
    "Referencia",
]

MASTER_HEADERS = [
    "Cantidad a la mano",
    "Cantidad pronosticada",
    "Costo",
    "Nombre",
    "Precio de venta",
    "Referencia interna",
    "Unidad de medida",
    "Costo de consignación",
]

MIXED_HEADERS = [
    "Proveedores/Proveedor/Referencia",
    "Proveedores/Proveedor",
    "Referencia interna",
    "Nombre",
    "Precio de venta",
    "Costo de consignación",
    "Cantidad a la mano",
    "Creado el",
    "Producto/Movimiento de stock",
    "Producto/Movimiento de stock/Cantidad",
    "Producto/Movimiento de stock/Fecha programada",
]


def _workbook(data: bytes) -> openpyxl.Workbook:
    return load_workbook(data)


def test_detect_import_type_contacts() -> None:
    assert mappings.detect_import_type(CONTACT_HEADERS) is ImportBatchType.CONTACTS


def test_detect_import_type_products() -> None:
    assert mappings.detect_import_type(MASTER_HEADERS) is ImportBatchType.PRODUCTS


def test_detect_import_type_mixed() -> None:
    assert mappings.detect_import_type(MIXED_HEADERS) is ImportBatchType.MIXED_PRODUCT_EXPORT


def test_detect_import_type_unknown() -> None:
    with pytest.raises(UnknownImportTypeError):
        mappings.detect_import_type(["Columna", "Desconocida"])


def test_missing_header_raises() -> None:
    with pytest.raises(MissingHeaderError):
        mappings.build_column_map(["Otra", "Cosa"], required=("Nombre completo",))


def test_parse_contacts() -> None:
    data = make_xlsx(
        CONTACT_HEADERS,
        [
            ["", "", "", "", "", "THEOBROMA SAC", "", "Perú", "", "", "PROVEEDOR", "THE6"],
            ["", "", "", "", "", "CHOCOLATES SAC", "", "Perú", "", "", "PROVEEDOR", "CHO6"],
        ],
    )
    contacts, errors, total = parse_contacts(_workbook(data))
    assert total == 2
    assert errors == []
    assert contacts[0].external_ref == "THE6"
    assert contacts[1].name == "CHOCOLATES SAC"


def test_parse_product_master_with_cost() -> None:
    data = make_xlsx(
        MASTER_HEADERS,
        [
            ["5", "5", "13.5596", "7756339000084 PT-03", "28", "THE60005", "U", "0"],
            ["0", "0", "19", "100% cacao", "32.5", "CHO60006", "U", "0"],
        ],
    )
    products, errors, total, duplicates = parse_product_master(_workbook(data))
    assert total == 2
    assert errors == []
    assert duplicates == []
    assert products[0].internal_reference == "THE60005"
    assert products[0].cost == decimal.Decimal("13.5596")
    assert products[0].sale_price == decimal.Decimal("28")


def test_parse_mixed_hierarchy_associates_movements() -> None:
    data = make_xlsx(
        MIXED_HEADERS,
        [
            ["THE6", "THEO", "THE60005", "CACAO", "28", "0", "5", "2026-02-06", "", "", ""],
            ["", "", "", "", "", "", "", "", "Stock>Inventory", "3", "2026-02-13 13:57:15"],
            ["CHO6", "CHOCO", "CHO60006", "cacao", "32.5", "0", "8", "2026-02-06", "", "", ""],
        ],
    )
    products, movements, errors, total, duplicates = parse_mixed_product_export(_workbook(data))
    assert total == 3
    assert len(products) == 2
    assert len(movements) == 1
    assert movements[0].product_internal_ref == "THE60005"
    assert movements[0].quantity == decimal.Decimal("3")
    assert products[0].supplier_ref == "THE6"
    assert errors == []


def test_parse_mixed_detects_duplicate_product() -> None:
    data = make_xlsx(
        MIXED_HEADERS,
        [
            ["THE6", "THEO", "THE60005", "CACAO", "28", "0", "5", "", "", "", ""],
            ["THE6", "THEO", "THE60005", "CACAO", "28", "0", "5", "", "", "", ""],
        ],
    )
    _products, _movements, _errors, _total, duplicates = parse_mixed_product_export(
        _workbook(data)
    )
    assert duplicates == ["THE60005"]


def test_invalid_decimal_is_flagged_not_zeroed() -> None:
    data = make_xlsx(
        MASTER_HEADERS,
        [["5", "5", "NO_NUMERICO", "Producto", "28", "REF1", "Unidades", "0"]],
    )
    products, errors, _total, _duplicates = parse_product_master(_workbook(data))
    assert products[0].cost is None
    assert any(e.error_code == "INVALID_DECIMAL" for e in errors)


def test_resolve_effective_cost_rules() -> None:
    assert resolve_effective_cost(decimal.Decimal("10"), decimal.Decimal("0")) == (
        decimal.Decimal("10"),
        CostSource.COST,
    )
    assert resolve_effective_cost(decimal.Decimal("0"), decimal.Decimal("8")) == (
        decimal.Decimal("8"),
        CostSource.CONSIGNMENT,
    )
    assert resolve_effective_cost(None, None) == (decimal.Decimal("0"), CostSource.ZERO)


def test_parse_datetime_aware() -> None:
    parsed, error = parse_datetime_aware("2026-02-13 13:57:15")
    assert error is None
    assert parsed is not None
    assert parsed.tzinfo is not None
    assert parsed == dt.datetime(2026, 2, 13, 13, 57, 15, tzinfo=parsed.tzinfo)


def test_parse_decimal_empty_vs_invalid() -> None:
    assert parse_decimal("") == (None, None)
    assert parse_decimal(None) == (None, None)
    assert parse_decimal("12.5") == (decimal.Decimal("12.5"), None)
    value, error = parse_decimal("abc")
    assert value is None
    assert error == "INVALID_DECIMAL"


def test_workbook_from_bytes_and_roundtrip() -> None:
    data = make_xlsx(["A", "B"], [[1, 2]])
    workbook = load_workbook(data)
    sheet = workbook[workbook.sheetnames[0]]
    assert sheet.max_row == 2
