"""Tests F010: branding configurable, motor global de documentos y cierre.

Cubre configuracion organizacional (A), generacion idempotente con
superacion (F/G) y el cierre oficial de campana. La identidad NUNCA vive
hardcodeada: todo sale de ``organization_settings``.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import Iterator
from typing import Any

import openpyxl
import pytest
from sqlalchemy import func, select

from app.db.session import SessionLocal
from app.documents.modules.inventory.audit_export import SHEET_NAMES
from app.documents.modules.inventory.erp_export import (
    CANONICAL_COLUMNS,
    CANONICAL_PROFILE_CODE,
    VENDOR_TEMPLATE_WARNING,
)
from app.models import (
    AuditEvent,
    DocumentExport,
    ExportProfile,
    InventoryCampaign,
    OrganizationSettings,
)
from app.models.enums import DocumentStatus
from tests.auth_helpers import override_auth
from tests.inventory_helpers import (
    PERMS_CLOSE,
    PERMS_DOCS_ALL,
    PERMS_EXPORTS,
    PERMS_EXPORTS_READ,
    PERMS_READ,
    PERMS_SYSTEM,
    as_user,
    cleanup_inventory_test_data,
    client,
)
from tests.test_count_engine import _ready_to_count
from tests.test_reconciliation import _audit_count, _campaign
from tests.test_valuation import _unwrap, _valued_campaign, _vcalculate, _vread

API = "/api/v1"
_PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
# Identidad F010: permisos de documentos + lectura/valorizacion de campana.
_DOCS_ID = PERMS_DOCS_ALL | {"inventory.read", "inventory.approve", "inventory.reconcile"}
_IDENTITY = {
    "legal_name": "Dedalo Retail S.A.C.",
    "tax_id": "20600000001",
    "tax_id_label": "RUC",
    "brand_color_primary": "#1F4E79",
    "brand_color_secondary": "#F5F7FA",
    "brand_color_accent": "#C00000",
    "footer_left": "Inventario Dedalo",
    "footer_right": "Documento oficial",
}


@pytest.fixture(autouse=True)
def _cleanup_after_documents_test() -> Iterator[None]:
    yield
    cleanup_inventory_test_data()


# ------------------------------- helpers ------------------------------------


def _settings_get() -> tuple[int, dict[str, Any]]:
    response = client.get(f"{API}/system/organization-settings")
    return response.status_code, _unwrap(response)


def _settings_patch(payload: dict[str, Any]) -> tuple[int, dict[str, Any]]:
    response = client.patch(f"{API}/system/organization-settings", json=payload)
    return response.status_code, _unwrap(response)


def _settings_logo(data: bytes, mime: str = "image/png") -> tuple[int, dict[str, Any]]:
    response = client.post(
        f"{API}/system/organization-settings/logo",
        files={"file": ("logo.png", data, mime)},
    )
    return response.status_code, _unwrap(response)


def _configure_identity(**extra: object) -> dict[str, Any]:
    status, body = _settings_patch({**_IDENTITY, **extra})
    assert status == 200, body
    return body


def _generate(
    campaign_id: str,
    document_type: str,
    *,
    expected_version: int | None = None,
    profile_code: str | None = None,
) -> tuple[int, dict[str, Any]]:
    payload: dict[str, Any] = {}
    if expected_version is not None:
        payload["expected_version"] = expected_version
    if profile_code is not None:
        payload["profile_code"] = profile_code
    response = client.post(
        f"{API}/inventory/campaigns/{campaign_id}/documents/{document_type}",
        json=payload,
    )
    return response.status_code, _unwrap(response)


def _campaign_documents(
    campaign_id: str, **params: str
) -> tuple[int, dict[str, Any]]:
    query = "&".join(f"{key}={value}" for key, value in params.items())
    suffix = f"?{query}" if query else ""
    response = client.get(f"{API}/inventory/campaigns/{campaign_id}/documents{suffix}")
    return response.status_code, _unwrap(response)


def _doc_detail(document_id: str) -> tuple[int, dict[str, Any]]:
    response = client.get(f"{API}/documents/{document_id}")
    return response.status_code, _unwrap(response)


def _doc_download(document_id: str) -> tuple[int, Any]:
    response = client.get(f"{API}/documents/{document_id}/download")
    return response.status_code, response


def _global_documents(**params: str) -> tuple[int, dict[str, Any]]:
    query = "&".join(f"{key}={value}" for key, value in params.items())
    suffix = f"?{query}" if query else ""
    response = client.get(f"{API}/documents{suffix}")
    return response.status_code, _unwrap(response)


def _close(campaign_id: str, expected_version: int | None = None) -> tuple[int, dict[str, Any]]:
    payload = {} if expected_version is None else {"expected_version": expected_version}
    response = client.post(f"{API}/inventory/campaigns/{campaign_id}/close", json=payload)
    return response.status_code, _unwrap(response)


def _valued_for_documents() -> dict[str, Any]:
    """Campana APPROVED con valorizacion F009 persistida."""
    ctx = _valued_campaign((("A", "5"),), [("A", 3)])
    as_user(_DOCS_ID, roles=("MANAGER",))
    version = _campaign(ctx["campaign_id"])["version"]
    status, body = _vcalculate(ctx["campaign_id"], version)
    assert status == 200, body
    as_user(_DOCS_ID)
    return ctx


def _audit_by_action(action: str) -> int:
    with SessionLocal() as db:
        return int(
            db.execute(
                select(func.count()).select_from(AuditEvent).where(AuditEvent.action == action)
            ).scalar_one()
        )


def _document_count(campaign_id: str) -> int:
    with SessionLocal() as db:
        return int(
            db.execute(
                select(func.count())
                .select_from(DocumentExport)
                .where(DocumentExport.entity_id == uuid.UUID(campaign_id))
            ).scalar_one()
        )


def _document_row(document_id: str) -> DocumentExport:
    with SessionLocal() as db:
        row = db.get(DocumentExport, uuid.UUID(document_id))
        assert row is not None
        db.expunge(row)
        return row


def _workbook(content: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)


def _cells(worksheet: Any) -> list[list[Any]]:
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


# ------------------------------- F010A: settings -----------------------------


def test_organization_settings_require_system_manage() -> None:
    as_user(PERMS_READ)
    status, _body = _settings_get()
    assert status == 403
    status, _body = _settings_patch({"legal_name": "Hack"})
    assert status == 403
    status, _body = _settings_logo(_PNG)
    assert status == 403


def test_settings_singleton_defaults_and_audit() -> None:
    as_user(PERMS_SYSTEM)
    status, body = _settings_get()
    assert status == 200, body
    assert body["is_configured"] is False
    assert body["legal_name"] is None
    assert body["document_number_prefix"] == "DOC"
    assert body["next_document_number"] == 1
    assert body["version"] == 1
    assert body["logo"] is None
    assert body["currency_style"] == "SYMBOL_BEFORE"
    assert body["date_format"] == "YYYY-MM-DD"
    assert _audit_by_action("ORGANIZATION_SETTINGS_CREATED") == 1
    second_status, second = _settings_get()
    assert second_status == 200
    assert second["id"] == body["id"]
    assert _audit_by_action("ORGANIZATION_SETTINGS_CREATED") == 1


def test_settings_patch_updates_identity_with_optimistic_version() -> None:
    as_user(PERMS_SYSTEM)
    status, body = _settings_patch(_IDENTITY)
    assert status == 200, body
    assert body["legal_name"] == "Dedalo Retail S.A.C."
    assert body["tax_id"] == "20600000001"
    assert body["brand_color_primary"] == "#1F4E79"
    assert body["version"] == 2
    assert body["is_configured"] is True
    assert _audit_by_action("ORGANIZATION_SETTINGS_UPDATED") == 1

    status, conflict = _settings_patch({**_IDENTITY, "legal_name": "Otra", "expected_version": 1})
    assert status == 409, conflict
    assert conflict["error"] == "VERSION_CONFLICT"

    status, stale = _settings_get()
    assert status == 200
    assert stale["legal_name"] == "Dedalo Retail S.A.C."

    status, empty = _settings_patch({})
    assert status == 422, empty
    assert empty["error"] == "ORGANIZATION_NO_CHANGES"


def test_settings_validation_rejects_bad_brand_values() -> None:
    as_user(PERMS_SYSTEM)
    status, body = _settings_patch({"brand_color_primary": "rojo"})
    assert status == 422, body
    assert body["error"] == "ORGANIZATION_INVALID_COLOR"
    status, body = _settings_patch({"date_format": "DD.MM.YYYY"})
    assert status == 422, body
    assert body["error"] == "ORGANIZATION_INVALID_FIELD"
    status, body = _settings_patch({"document_number_prefix": "DOC 1/2"})
    assert status == 422, body
    assert body["error"] == "ORGANIZATION_INVALID_FIELD"
    # El tipado del request rechaza campos no textos antes de llegar al servicio.
    response = client.patch(
        f"{API}/system/organization-settings", json={"legal_name": 42}
    )
    assert response.status_code == 422
    assert _audit_by_action("ORGANIZATION_SETTINGS_UPDATED") == 0


def test_logo_upload_validation_and_persistence() -> None:
    as_user(PERMS_SYSTEM)
    status, body = _settings_logo(b"")
    assert status == 422, body
    assert body["error"] == "LOGO_EMPTY_FILE"
    status, body = _settings_logo(b"hola", "text/plain")
    assert status == 415, body
    assert body["error"] == "LOGO_UNSUPPORTED_MEDIA"
    status, body = _settings_logo(b"not-an-image", "image/png")
    assert status == 422, body
    assert body["error"] == "LOGO_INVALID_CONTENT"

    status, body = _settings_logo(_PNG)
    assert status == 200, body
    assert body["logo"]["sha256"] and len(body["logo"]["sha256"]) == 64
    assert body["logo"]["media_type"] == "image/png"
    assert body["logo"]["path"].endswith(".png")
    assert body["version"] == 2
    assert _audit_by_action("ORGANIZATION_LOGO_UPDATED") == 1
    # La ruta del logo es generada por backend (hash), nunca del filename.
    assert "logo.png" not in body["logo"]["path"]

    status, big = _settings_logo(b"\x89PNG\r\n\x1a\n" + b"\x00" * (6 * 1024 * 1024))
    assert status == 413, big
    assert big["error"] == "LOGO_TOO_LARGE"


# ------------------------------- F010F: generacion ---------------------------


def test_generation_requires_configured_organization_identity() -> None:
    ctx = _valued_for_documents()
    as_user(PERMS_EXPORTS)
    status, body = _generate(ctx["campaign_id"], "management-report")
    assert status == 409, body
    assert body["error"] == "ORGANIZATION_SETTINGS_INCOMPLETE"
    assert _document_count(ctx["campaign_id"]) == 0


def test_management_report_is_generated_with_configurable_branding() -> None:
    ctx = _valued_for_documents()
    as_user(_DOCS_ID)
    _configure_identity(document_number_prefix="INV")

    status, body = _generate(ctx["campaign_id"], "management-report")
    assert status == 200, body
    assert body["already_generated"] is False
    assert body["status"] == "GENERATED"
    assert body["format"] == "PDF"
    assert body["file"]["content_type"] == "application/pdf"
    assert body["document_number"].startswith("INV-")
    assert body["document_number"] == "INV-000001"
    assert body["module"] == "INVENTORY"
    assert body["entity_type"] == "inventory_campaign"
    assert body["entity_id"] == ctx["campaign_id"]
    assert len(body["source_sha256"]) == 64
    assert len(body["branding_sha256"]) == 64
    assert body["file"]["size"] > 0

    status, download = _doc_download(body["id"])
    assert status == 200, download.text
    assert download.headers["content-type"] == "application/pdf"
    assert download.content.startswith(b"%PDF")
    assert len(download.content) == body["file"]["size"]

    # El dinero sale de F009: el snapshot replica la valorizacion persistida.
    status, valuation = _vread(ctx["campaign_id"])
    assert status == 200, valuation
    snapshot = body["source_snapshot"]
    assert snapshot["valuation"]["summary"] == valuation["summary"]
    assert snapshot["valuation"]["currency"] == valuation["currency"]
    assert snapshot["valuation"]["items_count"] == len(valuation["items"])
    assert (
        snapshot["campaign"]["valuation_source_sha256"]
        == valuation["campaign"]["valuation_source_sha256"]
    )
    assert snapshot["export_profile"] is None
    assert _audit_count("DOCUMENT_GENERATED", ctx["campaign_id"]) == 1


def test_generation_is_idempotent_and_supersedes_on_branding_change() -> None:
    ctx = _valued_for_documents()
    as_user(_DOCS_ID)
    _configure_identity()

    status, first = _generate(ctx["campaign_id"], "management-report")
    assert status == 200, first
    status, second = _generate(ctx["campaign_id"], "management-report")
    assert status == 200, second
    assert second["already_generated"] is True
    assert second["id"] == first["id"]
    assert _document_count(ctx["campaign_id"]) == 1
    assert _audit_count("DOCUMENT_GENERATED", ctx["campaign_id"]) == 1

    status, patched = _settings_patch({"legal_name": "Dedalo Retail E.I.R.L."})
    assert status == 200, patched
    status, third = _generate(ctx["campaign_id"], "management-report")
    assert status == 200, third
    assert third["already_generated"] is False
    assert third["id"] != first["id"]
    assert third["branding_sha256"] != first["branding_sha256"]
    assert third["document_number"].endswith("000002")

    old = _document_row(first["id"])
    assert old.status is DocumentStatus.SUPERSEDED
    assert old.superseded_by == uuid.UUID(third["id"])
    assert old.superseded_at is not None
    assert _audit_count("DOCUMENT_SUPERSEDED", ctx["campaign_id"]) == 1
    assert _audit_count("DOCUMENT_GENERATED", ctx["campaign_id"]) == 2


def test_generation_requires_approved_campaign_and_valuation() -> None:
    campaign_id, _batch_id, _operator_id = _ready_to_count()
    as_user(_DOCS_ID)
    status, body = _generate(campaign_id, "management-report")
    assert status == 409, body
    assert body["error"] == "CAMPAIGN_NOT_APPROVED"

    ctx = _valued_campaign((("A", "5"),), [("A", 3)])
    as_user(_DOCS_ID)
    _configure_identity()
    status, body = _generate(ctx["campaign_id"], "management-report")
    assert status == 409, body
    assert body["error"] == "VALUATION_NOT_CALCULATED"
    assert _document_count(ctx["campaign_id"]) == 0


def test_generation_honors_expected_version_conflict() -> None:
    ctx = _valued_for_documents()
    as_user(_DOCS_ID)
    _configure_identity()
    version = _campaign(ctx["campaign_id"])["version"]
    status, body = _generate(
        ctx["campaign_id"], "management-report", expected_version=version - 5
    )
    assert status == 409, body
    assert _document_count(ctx["campaign_id"]) == 0


def test_audit_export_builds_all_sheets_from_persisted_data() -> None:
    ctx = _valued_for_documents()
    as_user(_DOCS_ID)
    _configure_identity()

    status, body = _generate(ctx["campaign_id"], "audit-export")
    assert status == 200, body
    assert body["format"] == "XLSX"
    status, download = _doc_download(body["id"])
    assert status == 200, download.text

    workbook = _workbook(download.content)
    assert workbook.sheetnames == list(SHEET_NAMES)
    metadata = {
        row[0]: row[1] for row in _cells(workbook["Metadatos"]) if row and row[0]
    }
    assert metadata["document_number"] == body["document_number"]
    assert metadata["source_sha256"] == body["source_sha256"]
    assert metadata["template_version"] == body["template_version"]
    valuation_sheet = _cells(workbook["Valorizacion"])
    assert valuation_sheet
    assert valuation_sheet[0][:9] == [
        "Referencia",
        "Producto",
        "Esperado",
        "Fisico aprobado",
        "Faltante",
        "Excedente",
        "Danadas",
        "Costo unitario efectivo",
        "Moneda",
    ]
    data_rows = [row for row in valuation_sheet[1:] if row and row[0]]
    # El dinero de la hoja sale de F009, sin recalculo del renderer.
    status, valuation = _vread(ctx["campaign_id"])
    assert status == 200, valuation
    assert valuation["items"]
    for item in valuation["items"]:
        reference = item["product"]["internal_reference"]
        row = next(r for r in data_rows if r[0] == reference)
        assert float(row[7]) == float(item["effective_unit_cost"])
        assert float(row[9]) == float(item["missing_cost_value"])
        assert float(row[10]) == float(item["damage_cost_value"])
        assert float(row[11]) == float(item["surplus_cost_value"])


def test_erp_adjustment_uses_canonical_profile_with_vendor_warning() -> None:
    ctx = _valued_for_documents()
    as_user(_DOCS_ID)
    _configure_identity()

    status, body = _generate(ctx["campaign_id"], "erp-adjustment")
    assert status == 200, body
    assert VENDOR_TEMPLATE_WARNING in body["warnings"]
    status, download = _doc_download(body["id"])
    assert status == 200, download.text

    workbook = _workbook(download.content)
    header = [cell for cell in next(workbook.active.iter_rows(values_only=True)) if cell]
    assert header == list(CANONICAL_COLUMNS)

    with SessionLocal() as db:
        profile = db.execute(
            select(ExportProfile).where(ExportProfile.code == CANONICAL_PROFILE_CODE)
        ).scalar_one()
        assert profile.active is True
        assert profile.requires_vendor_template is False
        assert list(profile.column_order or []) == list(CANONICAL_COLUMNS)

    status, unknown = _generate(
        ctx["campaign_id"], "erp-adjustment", profile_code="VENDOR_XYZ"
    )
    assert status == 404, unknown
    assert unknown["error"] == "EXPORT_PROFILE_NOT_FOUND"


# ------------------------------- F010G: historial ----------------------------


def test_document_history_list_detail_and_permissions() -> None:
    ctx = _valued_for_documents()
    as_user(_DOCS_ID)
    _configure_identity()
    status, report = _generate(ctx["campaign_id"], "management-report")
    assert status == 200, report
    status, audit = _generate(ctx["campaign_id"], "audit-export")
    assert status == 200, audit

    status, listing = _global_documents(entity_id=ctx["campaign_id"])
    assert status == 200, listing
    assert listing["total"] == 2
    assert {item["document_type"] for item in listing["items"]} == {
        "management-report",
        "audit-export",
    }

    status, filtered = _global_documents(
        entity_id=ctx["campaign_id"], document_type="management-report"
    )
    assert status == 200, filtered
    assert filtered["total"] == 1

    status, filtered = _global_documents(
        entity_id=ctx["campaign_id"], status="GENERATED"
    )
    assert status == 200, filtered
    assert filtered["total"] == 2

    status, detail = _doc_detail(report["id"])
    assert status == 200, detail
    assert detail["id"] == report["id"]
    assert detail["document_number"] == report["document_number"]

    missing = str(uuid.uuid4())
    status, _body = _doc_detail(missing)
    assert status == 404
    status, _body = _doc_download(missing)
    assert status == 404

    # Permisos separados: listar/descargar usa exports.read.
    override_auth({"exports.create"})
    status, _body = _global_documents(entity_id=ctx["campaign_id"])
    assert status == 403
    status, _body = _doc_detail(report["id"])
    assert status == 403
    status, _body = _doc_download(report["id"])
    assert status == 403
    override_auth({"exports.read"})
    status, listing = _campaign_documents(ctx["campaign_id"])
    assert status == 200, listing
    assert listing["total"] == 2


def test_generation_requires_exports_create_permission() -> None:
    ctx = _valued_for_documents()
    as_user(PERMS_READ)
    status, _body = _generate(ctx["campaign_id"], "management-report")
    assert status == 403
    as_user(PERMS_EXPORTS_READ)
    status, _body = _generate(ctx["campaign_id"], "management-report")
    assert status == 403
    assert _document_count(ctx["campaign_id"]) == 0


# ------------------------------- F010F: cierre -------------------------------


def test_close_requires_all_three_official_documents() -> None:
    ctx = _valued_for_documents()
    as_user(_DOCS_ID | PERMS_CLOSE)
    _configure_identity()
    campaign_id = ctx["campaign_id"]

    status, body = _close(campaign_id)
    assert status == 409, body
    assert body["error"] == "DOCUMENTS_REQUIRED"
    assert body["missing_documents"] == [
        "management-report",
        "audit-export",
        "erp-adjustment",
    ]

    status, report = _generate(campaign_id, "management-report")
    assert status == 200, report
    status, body = _close(campaign_id)
    assert status == 409, body
    assert body["missing_documents"] == ["audit-export", "erp-adjustment"]

    status, _audit_doc = _generate(campaign_id, "audit-export")
    assert status == 200, _audit_doc
    status, _erp_doc = _generate(campaign_id, "erp-adjustment")
    assert status == 200, _erp_doc

    version = _campaign(campaign_id)["version"]
    status, body = _close(campaign_id, expected_version=version)
    assert status == 200, body
    assert body["already_closed"] is False
    assert body["status"] == "CLOSED"
    assert body["closed_at"]
    assert set(body["documents"]) == {
        "management-report",
        "audit-export",
        "erp-adjustment",
    }
    assert _campaign(campaign_id)["status"] == "CLOSED"
    assert _audit_count("CAMPAIGN_CLOSED", campaign_id) == 1

    # Idempotencia: cerrar de nuevo no duplica auditoria ni cambia estado.
    version = _campaign(campaign_id)["version"]
    status, again = _close(campaign_id, expected_version=version)
    assert status == 200, again
    assert again["already_closed"] is True
    assert _audit_count("CAMPAIGN_CLOSED", campaign_id) == 1


def test_closed_campaign_blocks_generation_but_keeps_history_and_download() -> None:
    ctx = _valued_for_documents()
    as_user(_DOCS_ID | PERMS_CLOSE)
    _configure_identity()
    campaign_id = ctx["campaign_id"]
    status, report = _generate(campaign_id, "management-report")
    assert status == 200, report
    status, _audit_doc = _generate(campaign_id, "audit-export")
    assert status == 200, _audit_doc
    status, _erp = _generate(campaign_id, "erp-adjustment")
    assert status == 200, _erp
    status, closed = _close(campaign_id)
    assert status == 200, closed
    assert closed["already_closed"] is False

    status, body = _generate(campaign_id, "management-report")
    assert status == 409, body
    assert body["error"] == "CAMPAIGN_CLOSED"

    status, listing = _campaign_documents(campaign_id)
    assert status == 200, listing
    assert listing["total"] == 3
    status, download = _doc_download(report["id"])
    assert status == 200, download.text
    assert download.content.startswith(b"%PDF")
    assert _document_count(campaign_id) == 3


def test_close_requires_valuation_and_inventory_close_permission() -> None:
    ctx = _valued_for_documents()
    campaign_id = ctx["campaign_id"]
    with SessionLocal() as db:
        campaign = db.get(InventoryCampaign, uuid.UUID(campaign_id))
        assert campaign is not None
        campaign.valuation_calculated_at = None
        db.commit()

    as_user(_DOCS_ID | PERMS_CLOSE)
    _configure_identity()
    status, body = _close(campaign_id)
    assert status == 409, body
    assert body["error"] == "VALUATION_NOT_CALCULATED"

    override_auth(PERMS_EXPORTS)
    status, body = _close(campaign_id)
    assert status == 403, body


def test_settings_affect_subsequent_documents_only() -> None:
    ctx = _valued_for_documents()
    as_user(_DOCS_ID)
    _configure_identity()
    status, first = _generate(ctx["campaign_id"], "management-report")
    assert status == 200, first

    status, patched = _settings_patch({"date_format": "DD/MM/YYYY"})
    assert status == 200, patched
    status, second = _generate(ctx["campaign_id"], "management-report")
    assert status == 200, second
    # El cambio de fecha es identidad: hash de branding distinto => documento nuevo.
    assert second["branding_sha256"] != first["branding_sha256"]
    assert second["source_sha256"] != first["source_sha256"]
    assert _document_count(ctx["campaign_id"]) == 2


def test_organization_settings_row_is_singleton() -> None:
    as_user(PERMS_SYSTEM)
    _configure_identity()
    with SessionLocal() as db:
        count = int(
            db.execute(select(func.count()).select_from(OrganizationSettings)).scalar_one()
        )
    assert count == 1
    status, body = _settings_get()
    assert status == 200, body
    assert body["legal_name"] == "Dedalo Retail S.A.C."
