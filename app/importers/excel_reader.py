"""Lectura segura de XLSX con openpyxl (read_only, sin escribir)."""

from __future__ import annotations

import io
from collections.abc import Iterator
from pathlib import Path

import openpyxl

from app.importers.base import InvalidFileError

Source = bytes | str | Path


def load_workbook(source: Source) -> openpyxl.Workbook:
    """Abre un XLSX desde bytes o ruta, en modo solo lectura."""
    try:
        if isinstance(source, (str, Path)):
            return openpyxl.load_workbook(source, read_only=True, data_only=True)
        return openpyxl.load_workbook(io.BytesIO(source), read_only=True, data_only=True)
    except Exception as exc:  # pragma: no cover - openpyxl lanza varias excepciones
        raise InvalidFileError(f"No se pudo leer el archivo XLSX: {exc}") from exc


def first_sheet(workbook: openpyxl.Workbook) -> openpyxl.worksheet.worksheet.Worksheet:
    return workbook[workbook.sheetnames[0]]


def detect_header(
    sheet: openpyxl.worksheet.worksheet.Worksheet, min_non_empty: int = 2
) -> tuple[int, list[str]]:
    """Devuelve (numero de fila 1-based, encabezados) de la primera fila con datos."""
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = ["" if cell is None else str(cell).strip() for cell in row]
        if sum(1 for value in values if value) >= min_non_empty:
            return row_number, values
    return 0, []


def iter_data_rows(
    sheet: openpyxl.worksheet.worksheet.Worksheet, header_row: int
) -> Iterator[tuple[int, list[object]]]:
    """Itera filas de datos (despues de la cabecera) como (nro fila, valores)."""
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row_number <= header_row:
            continue
        yield row_number, list(row)
