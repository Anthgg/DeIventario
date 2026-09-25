"""Utilidades XLSX compartidas: anti formula-injection y escritura segura (F010)."""

from __future__ import annotations

import datetime as dt
import decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.documents.engine.formatting import excel_safe


def coerce(value: Any) -> Any:
    """Convierte tipos de dominio a tipos que openpyxl sabe escribir."""
    if value is None:
        return None
    if isinstance(value, decimal.Decimal):
        return float(value)
    if isinstance(value, dt.datetime):
        # Excel has no timezone-aware date type; persist a real UTC date cell.
        return (
            value.astimezone(dt.UTC).replace(tzinfo=None)
            if value.tzinfo is not None
            else value
        )
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def set_cell(worksheet: Worksheet, row: int, column: int, value: Any) -> None:
    cell = worksheet.cell(row=row, column=column)
    cell.value = excel_safe(coerce(value))


def write_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
    *,
    widths: list[float] | None = None,
) -> Worksheet:
    worksheet = workbook.create_sheet(title=title[:31])
    for index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=index)
        cell.value = excel_safe(header)
        cell.font = Font(bold=True)
    for row_offset, row_values in enumerate(rows, start=2):
        for column, value in enumerate(row_values, start=1):
            set_cell(worksheet, row_offset, column, value)
    if widths is None:
        widths = [max(12.0, min(40.0, len(header) + 6)) for header in headers]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.freeze_panes = "A2"
    return worksheet


def to_bytes(workbook: Workbook) -> bytes:
    import io

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
